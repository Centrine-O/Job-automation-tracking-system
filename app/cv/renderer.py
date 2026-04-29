"""Render a tailored CV dict to DOCX and PDF matching Centrine's resume format."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.colors import black
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable

OUTPUT_DIR = Path("data/cvs")


def _add_hyperlink(paragraph, url: str, text: str, size: int = 10):
    """Add a clickable hyperlink run to a DOCX paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    run_el.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run_el.append(t)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _slug(s: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in s)[:40]


def render(cv: dict) -> dict:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    base = OUTPUT_DIR / f"cv_{cv['job_id']}_{_slug(cv['company'])}_{_slug(cv['job_title'])}"
    docx_path = str(base) + ".docx"
    pdf_path = str(base) + ".pdf"
    _render_docx(cv, docx_path)
    _render_pdf(cv, pdf_path)
    return {"docx": docx_path, "pdf": pdf_path}


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _add_border_bottom(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _render_docx(cv: dict, path: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        _add_border_bottom(p)

    def body(text, size=10):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text).font.size = Pt(size)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.add_run(text).font.size = Pt(10)

    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(cv["name"])
    run.bold = True
    run.font.size = Pt(16)

    # Location
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(cv["location"]).font.size = Pt(10)

    # Contact line with clickable LinkedIn / GitHub
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    contact_text = f"{cv['email']} | {cv['phone']}"
    if cv.get("linkedin") or cv.get("github"):
        contact_text += " | "
    run = p.add_run(contact_text)
    run.font.size = Pt(10)
    if cv.get("linkedin"):
        _add_hyperlink(p, cv["linkedin"], "LinkedIn")
        if cv.get("github"):
            p.add_run(" | ").font.size = Pt(10)
    if cv.get("github"):
        _add_hyperlink(p, cv["github"], "GitHub")

    # Profile Summary
    heading("PROFILE SUMMARY")
    body(cv.get("profile_summary", ""))

    # Skills
    heading("SKILLS")
    for cat, items in cv.get("skills", {}).items():
        if items:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            bold_run = p.add_run(f"{cat}: ")
            bold_run.bold = True
            bold_run.font.size = Pt(10)
            p.add_run(", ".join(items)).font.size = Pt(10)

    # Experience
    heading("PROFESSIONAL EXPERIENCE")
    for role in cv.get("experience", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(role["title"])
        run.bold = True
        run.font.size = Pt(10)

        meta_parts = [role["company"]]
        if role.get("location"):
            meta_parts.append(role["location"])
        meta_parts.append(f"{role['start_date']} – {role['end_date']}")

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(" | ".join(meta_parts)).font.size = Pt(10)

        for b in role.get("bullets", []):
            bullet(b)

    # Certifications
    heading("CERTIFICATIONS")
    for cert in cv.get("certifications", []):
        name = cert.get("name", str(cert))
        issuer = cert.get("issuer", "")
        bullet(f"{name} – {issuer}" if issuer else name)

    # Education
    heading("EDUCATION")
    for edu in cv.get("education", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"{edu.get('degree', '')}")
        run.bold = True
        run.font.size = Pt(10)
        body(edu.get("institution", ""))

    body("Referees: Available on request")
    doc.save(path)


# ── PDF ───────────────────────────────────────────────────────────────────────

def _render_pdf(cv: dict, path: str):
    doc = SimpleDocTemplate(
        path, pagesize=A4,
        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
        leftMargin=0.8 * inch, rightMargin=0.8 * inch,
    )

    S = {
        "name": ParagraphStyle("name", fontSize=16, fontName="Helvetica-Bold",
                               alignment=TA_CENTER, spaceAfter=6),
        "contact": ParagraphStyle("contact", fontSize=10, fontName="Helvetica",
                                  alignment=TA_CENTER, spaceAfter=10),
        "section": ParagraphStyle("section", fontSize=11, fontName="Helvetica-Bold",
                                  spaceBefore=8, spaceAfter=3),
        "body": ParagraphStyle("body", fontSize=10, fontName="Helvetica",
                               spaceAfter=4, leading=14),
        "bullet": ParagraphStyle("bullet", fontSize=10, fontName="Helvetica",
                                 leftIndent=14, spaceAfter=2, leading=13),
        "role_title": ParagraphStyle("role_title", fontSize=10,
                                     fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=1),
        "role_meta": ParagraphStyle("role_meta", fontSize=10, fontName="Helvetica",
                                    spaceAfter=2),
    }

    story = []

    story.append(Paragraph(cv["name"], S["name"]))
    story.append(Paragraph(cv["location"], S["contact"]))

    contact_parts = [cv["email"], cv["phone"]]
    if cv.get("linkedin"):
        contact_parts.append(f'<link href="{cv["linkedin"]}"><u><font color="#0563C1">LinkedIn</font></u></link>')
    if cv.get("github"):
        contact_parts.append(f'<link href="{cv["github"]}"><u><font color="#0563C1">GitHub</font></u></link>')
    story.append(Paragraph(" | ".join(contact_parts), S["contact"]))

    def section(title):
        story.append(HRFlowable(width="100%", thickness=0.5, color=black, spaceAfter=2))
        story.append(Paragraph(title, S["section"]))

    section("PROFILE SUMMARY")
    story.append(Paragraph(cv.get("profile_summary", ""), S["body"]))

    section("SKILLS")
    for cat, items in cv.get("skills", {}).items():
        if items:
            story.append(Paragraph(f"<b>{cat}:</b> {', '.join(items)}", S["body"]))

    section("PROFESSIONAL EXPERIENCE")
    for role in cv.get("experience", []):
        story.append(Paragraph(role["title"], S["role_title"]))
        meta_parts = [role["company"]]
        if role.get("location"):
            meta_parts.append(role["location"])
        meta_parts.append(f"{role['start_date']} – {role['end_date']}")
        story.append(Paragraph(" | ".join(meta_parts), S["role_meta"]))
        for b in role.get("bullets", []):
            story.append(Paragraph(f"• {b}", S["bullet"]))

    section("CERTIFICATIONS")
    for cert in cv.get("certifications", []):
        name = cert.get("name", str(cert))
        issuer = cert.get("issuer", "")
        story.append(Paragraph(f"• {name} – {issuer}" if issuer else f"• {name}", S["bullet"]))

    section("EDUCATION")
    for edu in cv.get("education", []):
        story.append(Paragraph(f"<b>{edu.get('degree','')}</b>", S["role_title"]))
        story.append(Paragraph(edu.get("institution", ""), S["role_meta"]))

    story.append(Paragraph("Referees: Available on request", S["body"]))

    doc.build(story)
